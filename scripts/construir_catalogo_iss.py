"""Construye db/seeds/iss2001_catalogo.csv: el catálogo de procedimientos del Manual ISS 2001.

Une dos fuentes:
  - db/seeds/iss2001_uvr.csv (procedimientos quirúrgicos con UVR, extraídos del PDF)
  - el Excel consolidado de las tablas del manual (todos los capítulos)

Uso: python3 scripts/construir_catalogo_iss.py <tablas_consolidadas.xlsx>
"""
import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
CODIGO = re.compile(r"[A-Z0-9]{6}")


def limpiar(descripcion: str) -> str:
    d = " ".join(str(descripcion).split())
    d = re.sub(r"\s*\(\d+\)", "", d)          # notas al pie: "(65)"
    d = re.sub(r"\s*\(\d+$", "", d)           # nota al pie cortada al final: "(19"
    d = re.sub(r"\s*\*+$", "", d)             # asteriscos de notas
    return d.strip(" ,;")


def main(ruta_excel: str) -> None:
    catalogo: dict[str, dict] = {}
    for r in csv.DictReader(open(RAIZ / "db/seeds/iss2001_uvr.csv", encoding="utf-8")):
        catalogo[r["codigo_cups"]] = {"codigo": r["codigo_cups"], "descripcion": limpiar(r["descripcion"]),
                                      "uvr": r["uvr"], "especialidad": r["especialidad"]}
    hoja = load_workbook(ruta_excel, read_only=True, data_only=True).active
    for codigo, _tabla, descripcion, _valor, *_ in list(hoja.iter_rows(values_only=True))[1:]:
        codigo = str(codigo or "").strip().upper()
        if not CODIGO.fullmatch(codigo) or not descripcion:
            continue
        descripcion = limpiar(descripcion)
        actual = catalogo.get(codigo)
        if actual is None:
            # Del Excel no se toma el valor: en capítulos no quirúrgicos está en pesos, no en UVR
            catalogo[codigo] = {"codigo": codigo, "descripcion": descripcion, "uvr": "", "especialidad": ""}
        elif len(descripcion) > len(actual["descripcion"]) and actual["descripcion"] in descripcion:
            actual["descripcion"] = descripcion  # la del PDF pudo quedar cortada por el formato
    salida = RAIZ / "db/seeds/iss2001_catalogo.csv"
    with open(salida, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["codigo", "descripcion", "uvr", "especialidad"])
        w.writeheader()
        for fila in sorted(catalogo.values(), key=lambda x: x["codigo"]):
            w.writerow(fila)
    print(f"{len(catalogo)} procedimientos -> {salida.relative_to(RAIZ)}")


if __name__ == "__main__":
    main(sys.argv[1])
