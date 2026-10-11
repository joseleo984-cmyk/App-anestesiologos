"""Exportación a Excel de las cirugías y del resumen por cirujano, clínica y aseguradora."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PESOS = '"$"#,##0'
ENCABEZADO = Font(bold=True, color="FFFFFF")
FONDO = PatternFill("solid", fgColor="0B6BCB")


def _hoja(ws, titulos: list[str], filas: list[list], columnas_pesos: set[int], anchos: list[int]) -> None:
    ws.append(titulos)
    for celda in ws[1]:
        celda.font, celda.fill = ENCABEZADO, FONDO
    for fila in filas:
        ws.append(fila)
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
        if i in columnas_pesos:
            for (celda,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                celda.number_format = PESOS
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def libro_cirugias(cirugias: list[dict], reporte: dict, glosas: dict | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cirugías"
    filas = []
    for c in cirugias:
        procs = c["procedimientos"] or []
        filas.append([
            c["fecha_procedimiento"], c["paciente"], c["tipo_documento"], c["numero_documento"],
            c["aseguradora"], c["tipo_aseguradora"] or "PARTICULAR", c["numero_poliza"], c["clinica"],
            c["cirujano"], ", ".join(p["codigo"] for p in procs), "; ".join(p["nombre"] for p in procs),
            c["estado"], float(c["valor_honorarios"]) if c["valor_honorarios"] is not None else None,
            c["fecha_pago"],
        ])
    _hoja(ws, ["Fecha", "Paciente", "Tipo doc.", "Documento", "Aseguradora", "Tipo", "Póliza",
               "Clínica", "Cirujano", "Códigos CUPS", "Procedimientos", "Estado", "Valor honorarios",
               "Fecha de pago"], filas, {13}, [12, 30, 9, 14, 24, 12, 14, 28, 28, 16, 45, 12, 16, 14])
    for (celda,) in ws.iter_rows(min_row=2, min_col=1, max_col=1):
        celda.number_format = "yyyy-mm-dd"

    for titulo, clave in (("Por cirujano", "por_cirujano"), ("Por clínica", "por_clinica"),
                          ("Por aseguradora", "por_aseguradora"), ("Por mes", "por_mes")):
        hoja = wb.create_sheet(titulo)
        _hoja(hoja, [titulo.removeprefix("Por ").capitalize(), "Cirugías", "Total", "Pagado", "Por cobrar",
                     "Sin valor"],
              [[r["nombre"], r["cirugias"], float(r["total"]), float(r["pagado"]), float(r["por_cobrar"]),
                r["sin_valor"]] for r in reporte[clave]],
              {3, 4, 5}, [34, 10, 16, 16, 16, 10])

    if glosas and glosas["lineas"]:
        hoja = wb.create_sheet("Glosas")
        _hoja(hoja, ["Fecha servicio", "Empresa", "Documento", "Paciente", "Factura", "Autorización",
                     "Total", "Glosas", "Descuentos", "Neto", "Fecha de proceso", "Pagador"],
              [[g["fecha_servicio"], g["empresa"], g["documento_paciente"], g["paciente"], g["factura"],
                g["autorizacion"], *(float(g[k]) if g[k] is not None else None
                                     for k in ("total", "glosas", "descuentos", "neto")),
                g["fecha_proceso"], g["pagador"]] for g in glosas["lineas"]],
              {7, 8, 9, 10}, [14, 30, 14, 30, 10, 16, 14, 14, 14, 14, 14, 30])

    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def libro_relacion(filas: list[dict], anestesiologo: str, mes: str, fecha) -> bytes:
    """Relación de órdenes con la misma estructura del formato en papel."""
    from openpyxl.styles import Alignment, Border, Side
    wb = Workbook()
    ws = wb.active
    ws.title = "Relación de órdenes"
    linea = Side(style="thin", color="000000")
    borde = Border(left=linea, right=linea, top=linea, bottom=linea)
    ws.append(["RELACIÓN DE ÓRDENES"])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([f"ANESTESIÓLOGO: {anestesiologo}", None, None, f"FECHA: {fecha:%d/%m/%Y}", None,
               f"MES A FACTURAR: {mes}"])
    ws.append(["VERIFICADO POR:"])
    ws.append([])
    titulos = ["#", "ENTIDAD", "PACIENTE", "CÉDULA", "PROCEDIMIENTO", "Nº ORDEN", "VALOR"]
    ws.append(titulos)
    for celda in ws[5]:
        celda.font, celda.fill, celda.border = ENCABEZADO, FONDO, borde
    for i, f in enumerate(filas, start=1):
        ws.append([i, f["entidad"], f["paciente"], f["cedula"], f["procedimiento"], f["numero_orden"],
                   float(f["valor"]) if f["valor"] is not None else None])
    total = len(filas) + 6
    ws.append([None] * 5 + ["TOTAL", f"=SUM(G6:G{total - 1})" if filas else 0])
    for fila in ws.iter_rows(min_row=6, max_row=total):
        for celda in fila:
            celda.border = borde
            celda.alignment = Alignment(vertical="center", wrap_text=celda.column == 5)
        fila[6].number_format = PESOS
    ws.cell(total, 6).font = ws.cell(total, 7).font = Font(bold=True)
    for letra, ancho in zip("ABCDEFG", [5, 18, 32, 14, 48, 24, 14]):
        ws.column_dimensions[letra].width = ancho
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "5:5"
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
