"""Bitácora de validación: los anestesiólogos comparan las cifras de la app con otra fuente
(SimTIVA, una bomba TCI, otra app o una publicación) y dejan constancia de cada prueba.

Las pruebas son de casos simulados (sin datos de pacientes) y las ven todos los usuarios:
así se sabe qué modelos ya están probados y por cuántas personas.
"""
import io

from openpyxl import Workbook
from psycopg.types.json import Jsonb

from app import config
from app.exportar import _hoja

REFERENCIAS = {"simtiva": "SimTIVA", "bomba": "Bomba TCI", "otra_app": "Otra app", "publicacion": "Publicación", "otro": "Otro"}
VEREDICTOS = {"coincide": "Coincide", "aceptable": "Diferencia aceptable", "no_coincide": "No coincide"}


def diferencia_pct(app, referencia) -> float | None:
    """Diferencia de la app frente a la referencia, en % de la referencia (igual que Tiva.diferenciaPct)."""
    if app is None or referencia is None:
        return None
    if referencia == 0:
        return 0.0 if app == 0 else None
    return (app - referencia) / abs(referencia) * 100


def mayor_diferencia(comparacion: list[dict]) -> float | None:
    """La diferencia más grande en valor absoluto entre los puntos comparados (con su signo)."""
    difs = [d for p in comparacion if (d := diferencia_pct(p.get("app"), p.get("referencia"))) is not None]
    return max(difs, key=abs) if difs else None


def crear(conn, anestesiologo_id: int, prueba) -> dict:
    fila = conn.execute(
        """INSERT INTO validacion_prueba (anestesiologo_id, modulo, elemento, caso, referencia, referencia_detalle,
                                          parametros, comparacion, veredicto, comentario, version_app)
           VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s) RETURNING prueba_id""",
        (anestesiologo_id, prueba.modulo, prueba.elemento, prueba.caso, prueba.referencia, prueba.referencia_detalle,
         Jsonb(prueba.parametros), Jsonb([p.model_dump() for p in prueba.comparacion]), prueba.veredicto,
         prueba.comentario, config.VERSION)).fetchone()
    conn.commit()
    return fila


def listar(conn, modulo: str, anestesiologo_id: int) -> list[dict]:
    filas = conn.execute(
        """SELECT v.prueba_id, v.anestesiologo_id, a.nombre AS autor, v.elemento, v.caso, v.referencia,
                  v.referencia_detalle, v.parametros, v.comparacion, v.veredicto, v.comentario, v.version_app, v.creado_en
             FROM validacion_prueba v JOIN anestesiologo a USING (anestesiologo_id)
            WHERE v.modulo = %s ORDER BY v.creado_en DESC, v.prueba_id DESC""", (modulo,)).fetchall()
    for f in filas:
        f["mia"] = f["anestesiologo_id"] == anestesiologo_id
        f["mayor_diferencia"] = mayor_diferencia(f["comparacion"])
    return filas


def resumen(pruebas: list[dict]) -> list[dict]:
    """Por elemento (modelo): cuántas pruebas, de cuántas personas, cómo salieron y la mayor diferencia."""
    grupos: dict[str, dict] = {}
    for p in pruebas:
        g = grupos.setdefault(p["elemento"], {"elemento": p["elemento"], "pruebas": 0, "autores": set(),
                                               "coincide": 0, "aceptable": 0, "no_coincide": 0, "mayor_diferencia": None})
        g["pruebas"] += 1
        g["autores"].add(p["anestesiologo_id"])
        g[p["veredicto"]] += 1
        d = p["mayor_diferencia"]
        if d is not None and (g["mayor_diferencia"] is None or abs(d) > abs(g["mayor_diferencia"])):
            g["mayor_diferencia"] = d
    return [{**g, "autores": len(g["autores"])} for g in sorted(grupos.values(), key=lambda g: -g["pruebas"])]


def borrar(conn, prueba_id: int, anestesiologo_id: int, es_admin: bool) -> bool | None:
    """Borra la prueba si es de quien la pide (o si es el administrador). None = no existe; False = no es suya."""
    fila = conn.execute("SELECT anestesiologo_id FROM validacion_prueba WHERE prueba_id = %s", (prueba_id,)).fetchone()
    if fila is None:
        return None
    if fila["anestesiologo_id"] != anestesiologo_id and not es_admin:
        return False
    conn.execute("DELETE FROM validacion_prueba WHERE prueba_id = %s", (prueba_id,))
    conn.commit()
    return True


def libro(pruebas: list[dict]) -> bytes:
    """Excel con una hoja de pruebas y otra con cada cifra comparada."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Pruebas"
    nombre = lambda p: p["parametros"].get("modelo_nombre") or p["elemento"]  # noqa: E731
    _hoja(ws, ["Nº", "Fecha", "Anestesiólogo", "Modelo", "Caso", "Comparado con", "Detalle", "Veredicto",
               "Mayor diferencia (%)", "Comentario", "Versión de la app"],
          [[p["prueba_id"], p["creado_en"].replace(tzinfo=None), p["autor"], nombre(p), p["caso"], REFERENCIAS[p["referencia"]],
            p["referencia_detalle"], VEREDICTOS[p["veredicto"]],
            round(p["mayor_diferencia"], 1) if p["mayor_diferencia"] is not None else None, p["comentario"], p["version_app"]]
           for p in pruebas], set(), [6, 18, 24, 20, 50, 16, 26, 20, 12, 50, 12])
    for (celda,) in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        celda.number_format = "yyyy-mm-dd hh:mm"
    filas = []
    for p in pruebas:
        for c in p["comparacion"]:
            d = diferencia_pct(c.get("app"), c.get("referencia"))
            filas.append([p["prueba_id"], nombre(p), c["etiqueta"], c["unidad"], c.get("app"), c.get("referencia"),
                          round(d, 1) if d is not None else None])
    _hoja(wb.create_sheet("Cifras"), ["Nº prueba", "Modelo", "Cifra", "Unidad", "App", "Referencia", "Diferencia (%)"],
          filas, set(), [10, 20, 34, 10, 12, 12, 14])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
