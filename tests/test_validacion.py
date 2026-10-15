"""Bitácora de validación de TIVA: registrar, ver, resumir, borrar y exportar pruebas (casos simulados, sin pacientes)."""
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import auth, main, validacion
from tests.test_api import cliente  # noqa: F401  (fixture)


def prueba(**cambios) -> dict:
    datos = {
        "modulo": "tiva", "elemento": "schnider", "caso": "Propofol · Schnider · 40 años, H, 70 kg, 170 cm · Ce 3 µg/mL",
        "referencia": "simtiva", "referencia_detalle": "SimTIVA (web)",
        "parametros": {"farmaco": "propofol", "modelo": "schnider", "modelo_nombre": "Schnider", "edad": 40, "peso": 70},
        "comparacion": [
            {"clave": "bolo", "etiqueta": "Bolo de inducción (1.er minuto)", "unidad": "mg", "app": 110, "referencia": 100},
            {"clave": "mlh_15", "etiqueta": "Velocidad a los 15 min", "unidad": "mL/h", "app": 50, "referencia": 52},
            {"clave": "despertar", "etiqueta": "Despertar (Ce < 1.2)", "unidad": "min", "app": 9.5, "referencia": None},
        ],
        "veredicto": "aceptable", "comentario": "El bolo sale 10 % más alto.",
    }
    return {**datos, **cambios}


def otro_cliente(conn) -> TestClient:
    conn.execute("INSERT INTO anestesiologo (nombre, email) VALUES ('Dr. Otro', 'otro@ejemplo.com')")
    conn.commit()
    c = TestClient(main.app)
    c.cookies.set(auth.COOKIE, auth.crear_sesion(2))
    return c


def test_registrar_ver_y_resumir(cliente, conn):
    c = cliente
    assert c.get("/api/validacion").json() == {"pruebas": [], "resumen": []}
    r = c.post("/api/validacion", json=prueba())
    assert r.status_code == 201
    otro = otro_cliente(conn)
    assert otro.post("/api/validacion", json=prueba(veredicto="coincide", comparacion=[
        {"clave": "bolo", "etiqueta": "Bolo", "unidad": "mg", "app": 100, "referencia": 125}])).status_code == 201
    assert otro.post("/api/validacion", json=prueba(elemento="minto", veredicto="no_coincide")).status_code == 201

    d = c.get("/api/validacion").json()
    # todos ven todas las pruebas, la más reciente primero, con su autor
    assert [(p["elemento"], p["autor"], p["mia"]) for p in d["pruebas"]] == [
        ("minto", "Dr. Otro", False), ("schnider", "Dr. Otro", False), ("schnider", "Dra. Prueba", True)]
    mia = d["pruebas"][2]
    assert mia["mayor_diferencia"] == 10 and mia["version_app"] and mia["comparacion"][2]["referencia"] is None
    schnider = next(g for g in d["resumen"] if g["elemento"] == "schnider")
    assert schnider == {"elemento": "schnider", "pruebas": 2, "autores": 2, "coincide": 1, "aceptable": 1,
                        "no_coincide": 0, "mayor_diferencia": -20}
    assert d["resumen"][0]["elemento"] == "schnider"   # el más probado primero


def test_validaciones_de_la_prueba(cliente):
    c = cliente
    sin_referencia = prueba(comparacion=[{"clave": "bolo", "etiqueta": "Bolo", "unidad": "mg", "app": 100, "referencia": None}])
    assert c.post("/api/validacion", json=sin_referencia).status_code == 422
    assert c.post("/api/validacion", json=prueba(modulo="regional")).status_code == 422
    assert c.post("/api/validacion", json=prueba(referencia="google")).status_code == 422
    assert c.post("/api/validacion", json=prueba(elemento="Schnider; drop")).status_code == 422
    assert c.post("/api/validacion", json=prueba(parametros={"x": "a" * 5000})).status_code == 422
    assert c.post("/api/validacion", json=prueba(comparacion=[])).status_code == 422


def test_solo_el_autor_o_el_admin_borran(cliente, conn):
    c = cliente
    pid = c.post("/api/validacion", json=prueba()).json()["prueba_id"]
    otro = otro_cliente(conn)
    assert otro.delete(f"/api/validacion/{pid}").status_code == 403
    pid2 = otro.post("/api/validacion", json=prueba()).json()["prueba_id"]
    assert c.delete(f"/api/validacion/{pid}").status_code == 204
    assert c.delete(f"/api/validacion/{pid}").status_code == 404
    assert c.delete(f"/api/validacion/{pid2}").status_code == 403
    conn.execute("UPDATE anestesiologo SET es_admin = TRUE WHERE anestesiologo_id = 1")
    conn.commit()
    assert c.delete(f"/api/validacion/{pid2}").status_code == 204
    assert c.get("/api/validacion").json()["pruebas"] == []


def test_excel_de_la_bitacora(cliente):
    c = cliente
    c.post("/api/validacion", json=prueba())
    r = c.get("/api/validacion.xlsx")
    assert r.status_code == 200 and "bitacora-tiva-" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    pruebas, cifras = wb["Pruebas"], wb["Cifras"]
    assert [x.value for x in pruebas[2]][2:9] == ["Dra. Prueba", "Schnider", prueba()["caso"], "SimTIVA", "SimTIVA (web)",
                                                  "Diferencia aceptable", 10]
    assert [[x.value for x in fila] for fila in cifras.iter_rows(min_row=2)] == [
        [1, "Schnider", "Bolo de inducción (1.er minuto)", "mg", 110, 100, 10],
        [1, "Schnider", "Velocidad a los 15 min", "mL/h", 50, 52, -3.8],
        [1, "Schnider", "Despertar (Ce < 1.2)", "min", 9.5, None, None]]


def test_diferencias():
    assert validacion.diferencia_pct(110, 100) == 10 and validacion.diferencia_pct(None, 1) is None
    assert validacion.diferencia_pct(0, 0) == 0 and validacion.diferencia_pct(1, 0) is None
    assert validacion.mayor_diferencia([{"app": 95, "referencia": 100}, {"app": 130, "referencia": 100}]) == 30
    assert validacion.mayor_diferencia([{"app": 1, "referencia": None}]) is None
