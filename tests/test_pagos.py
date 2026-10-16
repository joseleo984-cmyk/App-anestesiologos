"""Reporte de pagos: lectura, cruce con las cirugías y aplicación (datos ficticios)."""
import json

import pytest

from app import extraccion
from app.modelos import ReportePagoExtraido
from tests.test_api import CONFIRMAR, cliente  # noqa: F401  (fixture)
from tests.test_extraccion import ClienteFalso, foto

REPORTE = {
    "pagador": "ASOCIACION DE ANESTESIOLOGOS (FICTICIA)", "prestador": "DRA. PRUEBA",
    "fecha_proceso": "2026-10-07", "periodo_desde": "2026-10-01", "periodo_hasta": "2026-10-31",
    "total_neto": 7419742,
    "filas": [
        {"empresa": "SEGUROS UNO", "documento_paciente": "72000730", "paciente": "PACIENTE UNO",
         "fecha_servicio": "2026-08-14", "factura": "27111", "autorizacion": "A1", "total": 700405,
         "glosas": 0, "descuentos": 0, "neto": 700405, "es_paciente": True},
        {"empresa": "PREPAGADA DOS", "documento_paciente": "32.622.878", "paciente": "PACIENTE DOS",
         "fecha_servicio": "2026-07-27", "factura": "26898", "autorizacion": "B2", "total": 591976,
         "glosas": 0, "descuentos": 1039, "neto": 590937, "es_paciente": True},
        {"empresa": "IPS TRES", "documento_paciente": "979-60", "paciente": "HONORARIOS X",
         "fecha_servicio": "2026-08-01", "factura": "26982", "autorizacion": "C3", "total": 6000000,
         "glosas": 0, "descuentos": 0, "neto": 6000000, "es_paciente": False},
        {"empresa": "SEGUROS UNO", "documento_paciente": "555000", "paciente": "NO REGISTRADO",
         "fecha_servicio": "2026-08-20", "factura": "1", "autorizacion": "D4", "total": 128400,
         "glosas": 0, "descuentos": 0, "neto": 128400, "es_paciente": True},
    ],
}


def crear(cliente, documento, fecha, estado="PROGRAMADA"):
    datos = {**CONFIRMAR, "fecha_procedimiento": fecha, "estado": estado,
             "paciente": {**CONFIRMAR["paciente"], "numero_documento": documento}}
    r = cliente.post("/api/cirugias", json=datos)
    assert r.status_code == 201, r.text
    return r.json()["cirugia_id"]


@pytest.fixture
def con_reporte(cliente, monkeypatch):
    def poner(reporte):
        monkeypatch.setattr(extraccion, "extraer_reporte_pago",
                            lambda d, t: ReportePagoExtraido.model_validate(reporte))
    poner(REPORTE)
    cliente.poner_reporte = poner
    return cliente


def subir(cliente, contenido=b"%PDF-1.7 reporte"):
    return cliente.post("/api/pagos", files={"archivo": ("reporte.pdf", contenido, "application/pdf")})


def test_cruce_y_aplicacion(con_reporte):
    c = con_reporte
    uno = crear(c, "72000730", "2026-08-14")                    # misma fecha
    dos = crear(c, "32622878", "2026-07-25")                    # 2 días de diferencia
    pendiente = crear(c, "111222", "2026-09-01")                # no está en el reporte
    crear(c, "555000", "2026-08-20", estado="CANCELADA")        # cancelada: no se cruza

    r = subir(c)
    assert r.status_code == 201, r.text
    rep = r.json()
    assert rep["cuadra"] is True and rep["estado"] == "LEIDO"
    assert [f["resultado"] for f in rep["filas"]] == ["COINCIDE", "COINCIDE", "OTRO_CONCEPTO", "SIN_CIRUGIA"]
    assert [f["cirugia_id"] for f in rep["filas"][:2]] == [uno, dos]

    assert c.post(f"/api/pagos/{rep['reporte_id']}/aplicar").json() == {"marcadas": 2}
    d = c.get(f"/api/cirugias/{uno}").json()
    assert d["estado"] == "PAGADA" and d["valor_honorarios"] == 700405 and d["fecha_pago"] == "2026-10-07"
    assert c.get(f"/api/cirugias/{dos}").json()["valor_honorarios"] == 590937  # el neto, con descuento

    # Aplicar otra vez no cambia nada: lo ya pagado no se toca
    assert c.post(f"/api/pagos/{rep['reporte_id']}/aplicar").json() == {"marcadas": 0}
    pendientes = c.get("/api/cirugias?estado=PENDIENTE").json()
    assert [x["cirugia_id"] for x in pendientes] == [pendiente]
    assert c.get("/api/reportes").json()["totales"]["pendientes"] == 1
    assert c.get("/api/pagos").json()[0]["encontradas"] == 2


def test_un_pago_no_se_cuenta_dos_veces(con_reporte):
    c = con_reporte
    crear(c, "72000730", "2026-08-14")
    rep = subir(c).json()
    c.post(f"/api/pagos/{rep['reporte_id']}/aplicar")
    # Otro reporte que repite el mismo paciente y fecha: la cirugía ya fue pagada por el primero
    c.poner_reporte({**REPORTE, "filas": REPORTE["filas"][:1], "total_neto": 700405})
    otro = subir(c, b"%PDF-1.7 otro reporte").json()
    assert otro["filas"][0]["resultado"] == "SIN_CIRUGIA"


def test_dos_cirugias_del_mismo_paciente(con_reporte):
    c = con_reporte
    primera = crear(c, "72000730", "2026-08-01")
    segunda = crear(c, "72000730", "2026-08-14")
    fila = REPORTE["filas"][0]
    c.poner_reporte({**REPORTE, "total_neto": None,
                     "filas": [fila, {**fila, "fecha_servicio": "2026-08-02", "factura": "2"}]})
    rep = subir(c).json()
    assert [f["cirugia_id"] for f in rep["filas"]] == [segunda, primera]
    assert rep["cuadra"] is None


def test_fuera_de_la_ventana_no_se_cruza(con_reporte):
    c = con_reporte
    crear(c, "72000730", "2025-10-01")  # más de 300 días antes del servicio (2026-08-14)
    crear(c, "72000730", "2026-09-23")  # más de 30 días después del servicio
    assert subir(c).json()["filas"][0]["resultado"] == "SIN_CIRUGIA"


def test_reporte_que_no_cuadra(con_reporte):
    c = con_reporte
    c.poner_reporte({**REPORTE, "total_neto": 1})
    assert subir(c).json()["cuadra"] is False


def test_duplicado_y_descartar(con_reporte):
    c = con_reporte
    rep = subir(c).json()
    assert subir(c).status_code == 409
    assert c.delete(f"/api/pagos/{rep['reporte_id']}").status_code == 204
    rep = subir(c).json()  # tras descartarlo se puede volver a subir
    c.post(f"/api/pagos/{rep['reporte_id']}/aplicar")
    assert c.delete(f"/api/pagos/{rep['reporte_id']}").status_code == 409
    assert c.get(f"/api/pagos/{rep['reporte_id']}/archivo").content == b"%PDF-1.7 reporte"


def test_lectura_del_reporte_con_claude():
    cliente = ClienteFalso(texto=json.dumps(REPORTE))
    r = extraccion.extraer_reporte_pago(b"%PDF-1.7", "application/pdf", client=cliente)
    kw = cliente.llamadas[0]
    assert kw["output_config"]["format"]["schema"] is extraccion.ESQUEMA_PAGO
    assert kw["output_config"]["effort"] == "medium" and kw["system"] is extraccion.INSTRUCCIONES_PAGO
    assert r.filas[1].documento_paciente == "32622878" and r.cuadra is True


def test_volver_a_cruzar_tras_registrar_el_faltante(con_reporte):
    c = con_reporte
    uno = crear(c, "72000730", "2026-08-14")
    rep = subir(c).json()
    c.post(f"/api/pagos/{rep['reporte_id']}/aplicar")
    assert rep["filas"][3]["resultado"] == "SIN_CIRUGIA"
    # se registra el paciente que faltaba: se cruza al guardarlo y, como el reporte ya estaba aplicado, queda pagado
    faltante = crear(c, "555000", "2026-08-20")
    rep = c.post(f"/api/pagos/{rep['reporte_id']}/cruzar").json()
    assert rep["filas"][0]["cirugia_id"] == uno                # lo ya pagado se conserva
    assert rep["filas"][3]["cirugia_id"] == faltante and rep["filas"][3]["estado_cirugia"] == "PAGADA"
    assert c.post(f"/api/pagos/{rep['reporte_id']}/aplicar").json() == {"marcadas": 0}
    assert c.get(f"/api/cirugias/{faltante}").json()["valor_honorarios"] == 128400


def test_reporte_de_glosas(con_reporte):
    c = con_reporte
    c.poner_reporte({**REPORTE, "total_neto": None, "filas": [
        {**REPORTE["filas"][0], "glosas": 50000, "neto": 650405},
        {**REPORTE["filas"][1]},                                             # solo descuento
        {**REPORTE["filas"][3]},                                             # sin glosa ni descuento
        {**REPORTE["filas"][0], "empresa": "OTRA", "fecha_servicio": "2026-01-10", "glosas": 10000}]})
    subir(c)
    g = c.get("/api/glosas").json()
    assert g["totales"]["glosas"] == 60000 and g["totales"]["descuentos"] == 1039 and g["totales"]["lineas"] == 3
    assert g["por_empresa"][0]["nombre"] == "SEGUROS UNO"
    assert c.get("/api/glosas?desde=2026-07-01").json()["totales"]["lineas"] == 2  # por fecha del servicio
    from openpyxl import load_workbook
    import io
    wb = load_workbook(io.BytesIO(c.get("/api/cirugias.xlsx").content))
    assert wb["Glosas"].max_row == 4


def test_cirugia_registrada_despues_del_reporte_se_cruza_sola(con_reporte):
    c = con_reporte
    # el reporte se sube primero, sin cirugías registradas; se aplica (lo que había)
    rep = subir(c).json()
    assert [f["resultado"] for f in rep["filas"]] == ["SIN_CIRUGIA", "SIN_CIRUGIA", "OTRO_CONCEPTO", "SIN_CIRUGIA"]
    c.post(f"/api/pagos/{rep['reporte_id']}/aplicar")
    # luego se registran las cirugías: se cruzan al guardarlas y, como el reporte ya estaba aplicado, quedan pagadas
    uno = crear(c, "72000730", "2026-08-14")
    assert c.get(f"/api/cirugias/{uno}").json()["estado"] == "PAGADA"
    filas = c.get(f"/api/pagos/{rep['reporte_id']}").json()["filas"]
    assert filas[0]["cirugia_id"] == uno and filas[0]["resultado"] == "COINCIDE"


def test_documento_con_prefijo_ceros_y_puntos(con_reporte):
    c = con_reporte
    c.poner_reporte({**REPORTE, "filas": [{**REPORTE["filas"][0], "documento_paciente": "CC 0072.000.730"}]})
    uno = crear(c, "72000730", "2026-08-14")
    rep = subir(c).json()
    assert rep["filas"][0]["cirugia_id"] == uno


def test_volver_a_cruzar_todos(con_reporte):
    c = con_reporte
    rep = subir(c).json()  # reporte leído (sin aplicar) antes de registrar las cirugías
    # se borra el cruce (como si las cirugías se hubieran registrado antes de esta versión): «cruzar todos» lo rehace
    from app.repositorio import conectar
    uno = crear(c, "72000730", "2026-08-14")
    with conectar() as conn:
        conn.execute("UPDATE pago_detalle SET cirugia_id = NULL")
        conn.commit()
    r = c.post("/api/pagos/cruzar-todos").json()
    assert r == {"reportes": 1, "lineas": 3, "encontradas": 1, "nuevas": 1, "marcadas_pagadas": 0, "por_aplicar": 1}
    assert c.get(f"/api/pagos/{rep['reporte_id']}").json()["filas"][0]["cirugia_id"] == uno
    assert c.get(f"/api/cirugias/{uno}").json()["estado"] == "PROGRAMADA"  # leído, no aplicado: no se marca pagada


def test_pago_hasta_300_dias_despues_de_la_cirugia(con_reporte):
    c = con_reporte
    # la fecha de servicio del reporte es la de radicación: la cirugía fue 210 días antes
    c.poner_reporte({**REPORTE, "fecha_proceso": "2026-10-07", "filas": [
        {**REPORTE["filas"][0], "fecha_servicio": "2026-08-14"},
        {**REPORTE["filas"][1], "fecha_servicio": "2026-07-27"}]})
    vieja = crear(c, "72000730", "2026-01-16")      # 210 días antes: se cruza
    crear(c, "32622878", "2026-10-20")              # posterior al pago (7 de octubre): no puede ser esa
    filas = subir(c).json()["filas"]
    assert filas[0]["cirugia_id"] == vieja
    assert filas[1]["cirugia_id"] is None


def test_con_varias_cirugias_toma_la_mas_cercana(con_reporte):
    c = con_reporte
    c.poner_reporte({**REPORTE, "filas": [REPORTE["filas"][0]]})   # servicio 2026-08-14
    crear(c, "72000730", "2026-02-01")
    cercana = crear(c, "72000730", "2026-07-30")
    assert subir(c).json()["filas"][0]["cirugia_id"] == cercana


def test_varias_lineas_del_mismo_paciente_van_a_la_misma_cirugia(con_reporte):
    c = con_reporte
    # AP paga en dos líneas los dos procedimientos de la misma cirugía
    linea = REPORTE["filas"][0]
    c.poner_reporte({**REPORTE, "total_neto": None, "filas": [
        {**linea, "neto": 500000, "factura": "1"}, {**linea, "neto": 200405, "factura": "2"}]})
    uno = crear(c, "72000730", "2026-07-20")
    rep = subir(c).json()
    assert [f["cirugia_id"] for f in rep["filas"]] == [uno, uno]
    assert c.post(f"/api/pagos/{rep['reporte_id']}/aplicar").json() == {"marcadas": 1}
    assert c.get(f"/api/cirugias/{uno}").json()["valor_honorarios"] == 700405  # la suma de las dos líneas


def test_marcar_todas_las_encontradas(con_reporte):
    c = con_reporte
    rep = subir(c).json()                                   # leído, sin aplicar
    uno = crear(c, "72000730", "2026-04-01")                # se cruzan al registrarlas
    dos = crear(c, "32622878", "2026-03-01")
    assert c.post("/api/pagos/cruzar-todos").json()["por_aplicar"] == 2
    assert c.post("/api/pagos/aplicar-todos").json() == {"marcadas": 2}
    assert {c.get(f"/api/cirugias/{i}").json()["estado"] for i in (uno, dos)} == {"PAGADA"}
    assert c.get(f"/api/pagos/{rep['reporte_id']}").json()["estado"] == "APLICADO"
    assert c.post("/api/pagos/cruzar-todos").json()["por_aplicar"] == 0


# ------------------------------------------------ relación de AP (clínicas que pagan el total, ej. Santa Ana)
RELACION = {
    "numero": "0061527", "fecha": "2026-09-01", "mes_facturar": "2026-09", "clinica": None,
    "anestesiologo": "DRA. PRUEBA", "total": 508360,
    "filas": [
        {"entidad": "ENTIDAD UNO", "paciente": "PACIENTE UNO", "documento_paciente": "72.000.730",
         "procedimiento": "CIRUGIA X", "numero_orden": "OR-77", "fecha_servicio": None, "valor": 300000},
        {"entidad": "ENTIDAD DOS", "paciente": "PACIENTE DOS", "documento_paciente": "32622878",
         "procedimiento": "CIRUGIA Y", "numero_orden": None, "fecha_servicio": None, "valor": 208360},
        {"entidad": None, "paciente": None, "documento_paciente": None, "procedimiento": None,
         "numero_orden": None, "fecha_servicio": None, "valor": None},            # renglón vacío
    ],
}


def crear_con_orden(cliente, documento, fecha, orden):
    datos = {**CONFIRMAR, "fecha_procedimiento": fecha, "numero_orden": orden,
             "paciente": {**CONFIRMAR["paciente"], "numero_documento": documento}}
    r = cliente.post("/api/cirugias", json=datos)
    assert r.status_code == 201, r.text
    return r.json()["cirugia_id"]


@pytest.fixture
def con_relacion(cliente, monkeypatch):
    llamadas = []

    def leer(archivos):
        llamadas.append(len(archivos))
        return extraccion.relacion_a_reporte(RELACION)
    monkeypatch.setattr(extraccion, "extraer_relacion", leer)
    cliente.llamadas_relacion = llamadas
    return cliente


def subir_relacion(cliente, paginas=(b"%PDF-1.7 hoja 1", b"%PDF-1.7 hoja 2"), **campos):
    return cliente.post("/api/pagos/relacion", data=campos,
                        files=[("archivos", (f"hoja{i}.pdf", p, "application/pdf")) for i, p in enumerate(paginas)])


def test_relacion_de_ap_marca_pagadas_con_su_valor(con_relacion):
    c = con_relacion
    otra = crear_con_orden(c, "72000730", "2026-08-25", "OR-99")      # más reciente, pero otra orden
    uno = crear_con_orden(c, "72000730", "2026-08-10", "OR-77")       # la de la relación (mismo nº de orden)
    dos = crear(c, "32622878", "2026-08-20")
    crear(c, "32622878", "2026-10-05")                                # después de la fecha de pago: no

    r = subir_relacion(c, pagador="Clínica Santa Ana (ficticia)", fecha_pago="2026-09-30")
    assert r.status_code == 201, r.text
    rep = r.json()
    assert c.llamadas_relacion == [2]                                 # las dos hojas juntas
    assert rep["tipo"] == "RELACION" and rep["numero"] == "0061527" and rep["cuadra"] is True and rep["paginas"] == 2
    assert rep["pagador"] == "Clínica Santa Ana (ficticia)" and rep["fecha_proceso"] == "2026-09-30"
    assert rep["periodo_desde"] == "2026-09-01" and rep["periodo_hasta"] == "2026-09-30"
    assert len(rep["filas"]) == 2 and rep["filas"][0]["numero_orden"] == "OR-77"
    assert [f["cirugia_id"] for f in rep["filas"]] == [uno, dos]

    assert c.post(f"/api/pagos/{rep['reporte_id']}/aplicar").json() == {"marcadas": 2}
    d = c.get(f"/api/cirugias/{uno}").json()
    assert d["estado"] == "PAGADA" and d["valor_honorarios"] == 300000 and d["fecha_pago"] == "2026-09-30"
    assert c.get(f"/api/cirugias/{otra}").json()["estado"] != "PAGADA"
    assert c.get(f"/api/pagos/{rep['reporte_id']}/archivo?pagina=2").content == b"%PDF-1.7 hoja 2"
    assert c.get(f"/api/pagos/{rep['reporte_id']}/archivo?pagina=3").status_code == 404
    assert c.get("/api/pagos").json()[0]["tipo"] == "RELACION"

    # Repetida: se avisa
    assert subir_relacion(c, pagador="x").status_code == 409


def test_relacion_no_cuenta_como_ingreso(con_relacion, conn):
    from app import finanzas
    assert subir_relacion(con_relacion).status_code == 201
    anest = conn.execute("SELECT anestesiologo_id FROM reporte_pago").fetchone()["anestesiologo_id"]
    assert finanzas.ingresos_por_mes(conn, anest) == []   # el pago global ya viene en el reporte de pagos


def test_lectura_de_la_relacion_con_claude():
    cliente = ClienteFalso(texto=json.dumps(RELACION))
    r = extraccion.extraer_relacion([(b"%PDF-1.7", "application/pdf"), (b"%PDF-1.7", "application/pdf")],
                                    client=cliente)
    kw = cliente.llamadas[0]
    assert kw["output_config"]["format"]["schema"] is extraccion.ESQUEMA_RELACION
    assert r.numero == "0061527" and r.fecha_proceso.isoformat() == "2026-09-01" and r.cuadra is True
    assert [f.documento_paciente for f in r.filas] == ["72000730", "32622878"]
    assert r.filas[0].neto == 300000 and r.filas[0].numero_orden == "OR-77"


# Radicado escrito a mano: sin valor por línea, solo el total al pie (datos ficticios)
RADICADO = {
    "numero": None, "fecha": "2026-01-15", "mes_facturar": None, "clinica": "Santa Ana",
    "anestesiologo": "DRA. PRUEBA", "total": 8957666,
    "filas": [
        {"entidad": "EPS FICTICIA", "paciente": "PACIENTE UNO", "documento_paciente": "72000730",
         "procedimiento": "CESAREA", "numero_orden": None, "fecha_servicio": None, "valor": None},
        {"entidad": None, "paciente": "PACIENTE DOS", "documento_paciente": "32622878",
         "procedimiento": "SEDACION PARA TAC", "numero_orden": None, "fecha_servicio": None, "valor": None},
        {"entidad": None, "paciente": "PACIENTE TRES", "documento_paciente": "55555555",
         "procedimiento": "COLGAJO", "numero_orden": None, "fecha_servicio": None, "valor": None},
    ],
}


@pytest.fixture
def con_radicado(cliente, monkeypatch):
    llamadas = []

    def leer(archivos, tipo="relacion"):
        llamadas.append(tipo)
        return extraccion.relacion_a_reporte(RADICADO)
    monkeypatch.setattr(extraccion, "extraer_relacion", leer)
    cliente.llamadas_radicado = llamadas
    return cliente


def subir_radicado(cliente, **campos):
    return cliente.post("/api/pagos/radicado", data=campos,
                        files=[("archivos", ("radicado.jpg", foto(), "image/jpeg"))])


def test_radicado_marca_facturadas_y_luego_pagadas_sin_borrar_el_valor(con_radicado):
    c = con_radicado
    datos = {**CONFIRMAR, "fecha_procedimiento": "2026-01-05", "valor_honorarios": 850000,
             "paciente": {**CONFIRMAR["paciente"], "numero_documento": "72000730"}}
    uno = c.post("/api/cirugias", json=datos).json()["cirugia_id"]
    dos = crear(c, "32622878", "2026-01-10")

    r = subir_radicado(c, pagador="Clínica Santa Ana (ficticia)")
    assert r.status_code == 201, r.text
    rep = r.json()
    assert c.llamadas_radicado == ["radicado"]
    assert rep["tipo"] == "RADICADO" and rep["total_neto"] == 8957666 and rep["fecha_proceso"] == "2026-01-15"
    assert rep["cuadra"] is None                       # las líneas no traen valor: no se compara con el total
    assert [f["resultado"] for f in rep["filas"]] == ["COINCIDE", "COINCIDE", "SIN_CIRUGIA"]

    # Un radicado no es un pago: no aparece como «pago encontrado» ni lo aplica «aplicar todos»
    inicio = c.get("/api/inicio").json()
    assert inicio["por_aplicar"] == 0
    assert inicio["radicados"][0]["total_neto"] == 8957666 and inicio["radicados"][0]["estado"] == "LEIDO"
    assert c.post("/api/pagos/aplicar-todos").json() == {"marcadas": 0}

    assert c.post(f"/api/pagos/{rep['reporte_id']}/aplicar").json() == {"marcadas": 2}
    d = c.get(f"/api/cirugias/{uno}").json()
    assert d["estado"] == "FACTURADA" and d["valor_honorarios"] == 850000 and d["fecha_pago"] is None
    pasos = c.get(f"/api/cirugias/{uno}/cobro").json()["pasos"]
    assert pasos[1]["hecho"] and pasos[1]["detalle"] == "Radicada en Clínica Santa Ana (ficticia)"
    assert c.get("/api/inicio").json()["radicados"][0]["cirugias"] == 2

    # Santa Ana paga el total: todas quedan pagadas con esa fecha, y el valor de cada cirugía se conserva
    r = c.post(f"/api/pagos/{rep['reporte_id']}/pagar", json={"fecha_pago": "2026-03-20"})
    assert r.json() == {"marcadas": 2}
    d = c.get(f"/api/cirugias/{uno}").json()
    assert d["estado"] == "PAGADA" and d["valor_honorarios"] == 850000 and d["fecha_pago"] == "2026-03-20"
    assert c.get(f"/api/cirugias/{dos}").json()["estado"] == "PAGADA"
    rep = c.get(f"/api/pagos/{rep['reporte_id']}").json()
    assert rep["estado"] == "PAGADO" and rep["pagado_en"] == "2026-03-20" and rep["valor_pagado"] == 8957666
    assert {f["resultado"] for f in rep["filas"]} == {"COINCIDE", "SIN_CIRUGIA"}
    assert c.get("/api/inicio").json()["radicados"] == []

    # La cirugía que faltaba se registra después: queda pagada de una vez con la fecha del pago
    tres = crear(c, "55555555", "2026-01-12")
    d = c.get(f"/api/cirugias/{tres}").json()
    assert d["estado"] == "PAGADA" and d["fecha_pago"] == "2026-03-20"


def test_solo_se_paga_un_radicado(con_relacion):
    rep = subir_relacion(con_relacion).json()
    r = con_relacion.post(f"/api/pagos/{rep['reporte_id']}/pagar", json={"fecha_pago": "2026-03-20"})
    assert r.status_code == 404


def test_aplicar_una_linea_sin_valor_no_borra_el_valor_de_la_cirugia(con_relacion, monkeypatch):
    c = con_relacion
    datos = {**CONFIRMAR, "fecha_procedimiento": "2026-08-20", "valor_honorarios": 400000,
             "paciente": {**CONFIRMAR["paciente"], "numero_documento": "32622878"}}
    cid = c.post("/api/cirugias", json=datos).json()["cirugia_id"]
    sin_valor = {**RELACION, "filas": [{**RELACION["filas"][1], "valor": None}]}
    monkeypatch.setattr(extraccion, "extraer_relacion", lambda archivos: extraccion.relacion_a_reporte(sin_valor))
    rep = subir_relacion(c, fecha_pago="2026-09-30").json()
    assert c.post(f"/api/pagos/{rep['reporte_id']}/aplicar").json() == {"marcadas": 1}
    d = c.get(f"/api/cirugias/{cid}").json()
    assert d["estado"] == "PAGADA" and d["valor_honorarios"] == 400000


def test_crear_desde_el_radicado_las_cirugias_que_faltan(con_radicado, monkeypatch):
    from app import radicado
    c = con_radicado
    dos = crear(c, "32622878", "2026-01-10")
    # dos líneas del mismo paciente (ficticio) son una sola cirugía; la entidad vacía queda por confirmar
    filas = [*RADICADO["filas"], {**RADICADO["filas"][0], "procedimiento": "POMEROY"},
             {"entidad": "SEGUROS FICTICIOS", "paciente": "ANA MARIA PRUEBA GOMEZ", "documento_paciente": "1.043.000.111",
              "procedimiento": "CESAREA + POMEROY", "numero_orden": None, "fecha_servicio": None, "valor": None}]
    monkeypatch.setattr(extraccion, "extraer_relacion",
                        lambda archivos, tipo="relacion": extraccion.relacion_a_reporte({**RADICADO, "filas": filas}))
    rep = subir_radicado(c, pagador="Clínica Santa Ana (ficticia)").json()
    assert c.post(f"/api/pagos/{rep['reporte_id']}/aplicar").json() == {"marcadas": 1}

    r = c.post(f"/api/pagos/{rep['reporte_id']}/crear-faltantes")
    assert r.status_code == 201, r.text
    cuerpo = r.json()
    assert cuerpo["creadas"] == 3
    assert {f["resultado"] for f in cuerpo["radicado"]["filas"]} == {"COINCIDE"}
    ids = {f["documento_paciente"]: f["cirugia_id"] for f in cuerpo["radicado"]["filas"]}
    assert ids["32622878"] == dos
    nueva = c.get(f"/api/cirugias/{ids['1043000111']}").json()
    assert nueva["estado"] == "FACTURADA"                              # el radicado ya estaba aplicado
    assert nueva["fecha_procedimiento"] == "2026-01-15" and nueva["valor_honorarios"] is None
    assert nueva["paciente"]["nombres"] == "ANA MARIA" and nueva["paciente"]["apellidos"] == "PRUEBA GOMEZ"
    assert nueva["aseguradora"]["tipo"] == "POLIZA"
    assert [p["nombre"] for p in nueva["procedimientos"]] == ["CESAREA", "POMEROY"]
    assert "radicado" in nueva["observaciones"]
    uno = c.get(f"/api/cirugias/{ids['72000730']}").json()           # dos líneas: una cirugía, dos procedimientos
    assert len(uno["procedimientos"]) == 2
    tres = c.get(f"/api/cirugias/{ids['55555555']}").json()
    assert tres["aseguradora"]["nombre"] == radicado.ENTIDAD_POR_CONFIRMAR

    # Repetirlo no duplica nada; al pagar el radicado, las creadas también quedan pagadas
    assert c.post(f"/api/pagos/{rep['reporte_id']}/crear-faltantes").json()["creadas"] == 0
    assert c.post(f"/api/pagos/{rep['reporte_id']}/pagar", json={"fecha_pago": "2026-03-20"}).json() == {"marcadas": 4}


def test_separar_nombre():
    from app.radicado import separar_nombre
    assert separar_nombre("Ana Pérez") == ("Ana", "Pérez")
    assert separar_nombre("Ana Pérez Gómez") == ("Ana", "Pérez Gómez")
    assert separar_nombre("Ana María Pérez Gómez") == ("Ana María", "Pérez Gómez")
    assert separar_nombre("Ana") == ("Ana", "SIN APELLIDO")


def test_pago_global_de_santa_ana_en_el_reporte_de_ap_paga_el_radicado(con_reporte, con_radicado, conn):
    c = con_reporte
    conn.execute("UPDATE anestesiologo SET nombre = 'Dra. Ana Ficticia López' WHERE anestesiologo_id = 1")
    conn.commit()
    uno = crear(c, "72000730", "2026-01-05")
    radicado = subir_radicado(c, pagador="Clínica Santa Ana (ficticia)").json()
    c.post(f"/api/pagos/{radicado['reporte_id']}/aplicar")

    # En la línea de Santa Ana viene la anestesióloga (nombre y cédula), no un paciente
    santa_ana = {"empresa": "CLINICA SANTA ANA (FICTICIA)", "documento_paciente": "99887766", "paciente": "FICTICIA LOPEZ",
                 "fecha_servicio": "2025-12-15", "factura": "9", "autorizacion": None, "total": 847275, "glosas": 0,
                 "descuentos": 42364, "neto": 804911, "es_paciente": True}
    otra = {**santa_ana, "paciente": "ANA F", "neto": 1000, "total": 1000, "descuentos": 0}
    c.poner_reporte({**REPORTE, "total_neto": None, "filas": [REPORTE["filas"][0], santa_ana, otra]})
    rep = subir(c).json()
    assert [f["resultado"] for f in rep["filas"]] == ["SIN_CIRUGIA", "PAGO_GLOBAL", "PAGO_GLOBAL"]
    assert [r["reporte_id"] for r in rep["radicados_por_cobrar"]] == [radicado["reporte_id"]]
    assert c.get("/api/inicio").json()["por_aplicar"] == 0

    linea = rep["filas"][1]["detalle_id"]
    r = c.post(f"/api/pagos/lineas/{linea}/radicado", json={"radicado_id": radicado["reporte_id"]})
    assert r.status_code == 200, r.text
    assert r.json()["marcadas"] == 1
    assert r.json()["reporte"]["filas"][1]["radicado_id"] == radicado["reporte_id"]
    d = c.get(f"/api/cirugias/{uno}").json()
    assert d["estado"] == "PAGADA" and d["fecha_pago"] == REPORTE["fecha_proceso"]
    pagado = c.get(f"/api/pagos/{radicado['reporte_id']}").json()
    assert pagado["estado"] == "PAGADO" and pagado["valor_pagado"] == 804911

    # Una línea de paciente no se puede usar como pago de un radicado
    paciente = rep["filas"][0]["detalle_id"]
    assert c.post(f"/api/pagos/lineas/{paciente}/radicado", json={"radicado_id": radicado["reporte_id"]}).status_code == 404


def crear_paciente(cliente, documento, fecha, nombres, apellidos):
    datos = {**CONFIRMAR, "fecha_procedimiento": fecha,
             "paciente": {"tipo_documento": "CC", "numero_documento": documento, "nombres": nombres, "apellidos": apellidos}}
    r = cliente.post("/api/cirugias", json=datos)
    assert r.status_code == 201, r.text
    return r.json()["cirugia_id"]


def test_paciente_guardado_con_la_cedula_mal_se_corrige_desde_el_radicado(con_radicado):
    c = con_radicado
    # «Paciente Tres» (ficticio) quedó guardado con un dígito mal en la cédula, en dos cirugías
    tres = crear_paciente(c, "55555556", "2026-01-12", "Paciente", "Tres")
    antes = crear_paciente(c, "55555556", "2025-11-02", "Paciente", "Tres")
    crear_paciente(c, "99999999", "2026-01-11", "Otra", "Persona")      # ni nombre ni cédula parecidos
    rep = subir_radicado(c, pagador="Clínica Santa Ana (ficticia)").json()
    c.post(f"/api/pagos/{rep['reporte_id']}/aplicar")
    linea = next(f for f in rep["filas"] if f["documento_paciente"] == "55555555")
    assert linea["resultado"] == "SIN_CIRUGIA"

    cand = c.get(f"/api/pagos/lineas/{linea['detalle_id']}/candidatas").json()["candidatas"]
    assert [x["cirugia_id"] for x in cand] == [tres, antes]                # la de fecha más cercana primero
    assert cand[0]["cedula_parecida"] is True and cand[0]["parecido"] == 100 and cand[0]["paciente"] == "Paciente Tres"

    # la cédula del radicado es la buena: se corrige la ficha del paciente (sus dos cirugías)
    r = c.post(f"/api/pagos/lineas/{linea['detalle_id']}/vincular", json={"cirugia_id": tres, "cedula_correcta": "radicado"})
    assert r.status_code == 200, r.text
    fila = next(f for f in r.json()["filas"] if f["detalle_id"] == linea["detalle_id"])
    assert fila["resultado"] == "COINCIDE" and fila["cirugia_id"] == tres
    d = c.get(f"/api/cirugias/{tres}").json()
    assert d["paciente"]["numero_documento"] == "55555555" and d["estado"] == "FACTURADA"   # el radicado ya estaba aplicado
    assert c.get(f"/api/cirugias/{antes}").json()["paciente"]["numero_documento"] == "55555555"
    # volver a cruzar no la pierde
    filas = c.post(f"/api/pagos/{rep['reporte_id']}/cruzar").json()["filas"]
    assert next(f for f in filas if f["detalle_id"] == linea["detalle_id"])["cirugia_id"] == tres


def test_cedula_mal_leida_en_el_radicado_se_corrige_en_la_linea(con_radicado, monkeypatch):
    c = con_radicado
    uno = crear_paciente(c, "72000730", "2026-01-05", "Paciente", "Uno")
    filas = [{**RADICADO["filas"][0], "documento_paciente": "72000780"}]   # la hoja a mano se leyó con un dígito mal
    monkeypatch.setattr(extraccion, "extraer_relacion",
                        lambda archivos, tipo="relacion": extraccion.relacion_a_reporte({**RADICADO, "filas": filas}))
    rep = subir_radicado(c, pagador="Clínica Santa Ana (ficticia)").json()
    linea = rep["filas"][0]
    assert linea["resultado"] == "SIN_CIRUGIA"
    assert c.get(f"/api/pagos/lineas/{linea['detalle_id']}/candidatas").json()["candidatas"][0]["cirugia_id"] == uno
    r = c.post(f"/api/pagos/lineas/{linea['detalle_id']}/vincular", json={"cirugia_id": uno, "cedula_correcta": "app"})
    fila = r.json()["filas"][0]
    assert fila["resultado"] == "COINCIDE" and fila["documento_paciente"] == "72000730"
    assert c.get(f"/api/cirugias/{uno}").json()["paciente"]["numero_documento"] == "72000730"
    # validaciones
    assert c.post(f"/api/pagos/lineas/{linea['detalle_id']}/vincular", json={"cirugia_id": uno, "cedula_correcta": "x"}).status_code == 422
    assert c.post(f"/api/pagos/lineas/{linea['detalle_id']}/vincular", json={"cirugia_id": 999, "cedula_correcta": "app"}).status_code == 404
    assert c.get("/api/pagos/lineas/999/candidatas").status_code == 404


def test_buscar_paciente_en_los_pagos(con_reporte, conn):
    c = con_reporte
    uno = crear_paciente(c, "72000730", "2026-08-14", "Paciente", "Uno")
    crear_paciente(c, "1.043.000.111", "2026-08-01", "Ana María", "Pérez Gómez")
    rep = subir(c).json()
    # por cédula, con o sin puntos y ceros: la cirugía y la línea del reporte que la pagó
    for q in ("72000730", "72.000.730", "0072000730", "2000730"):
        r = c.get("/api/pagos/buscar", params={"q": q}).json()
        assert [x["cirugia_id"] for x in r["cirugias"]] == [uno], q
        assert [(x["reporte_id"], x["cirugia_id"], x["neto"]) for x in r["lineas"]] == [(rep["reporte_id"], uno, 700405)], q
    # por nombre: sin tildes, en cualquier orden, en minúsculas
    r = c.get("/api/pagos/buscar", params={"q": "gomez ana"}).json()
    assert [x["paciente"] for x in r["cirugias"]] == ["Ana María Pérez Gómez"] and r["lineas"] == []
    r = c.get("/api/pagos/buscar", params={"q": "registrado"}).json()       # solo en el reporte
    assert r["cirugias"] == [] and [x["documento_paciente"] for x in r["lineas"]] == ["555000"]
    assert r["lineas"][0]["tipo"] == "PAGOS" and r["lineas"][0]["cirugia_id"] is None
    # los honorarios globales no son pacientes; lo de otro anestesiólogo no se ve
    assert c.get("/api/pagos/buscar", params={"q": "honorarios"}).json() == {"cirugias": [], "lineas": []}
    conn.execute("INSERT INTO anestesiologo (nombre, email) VALUES ('Dr. Otro', 'otro@ejemplo.com')")
    conn.commit()
    from fastapi.testclient import TestClient
    from app import auth, main
    otro = TestClient(main.app)
    otro.cookies.set(auth.COOKIE, auth.crear_sesion(2))
    assert otro.get("/api/pagos/buscar", params={"q": "72000730"}).json() == {"cirugias": [], "lineas": []}
    assert c.get("/api/pagos/buscar", params={"q": "a"}).status_code == 422
    assert c.get("/api/pagos/buscar", params={"q": "--"}).json() == {"cirugias": [], "lineas": []}


def test_radicado_con_sus_cirugias_pagadas_por_otro_lado_no_queda_por_cobrar(con_radicado):
    c = con_radicado
    uno = crear(c, "72000730", "2026-01-05")
    dos = crear(c, "32622878", "2026-01-10")
    rep = subir_radicado(c, pagador="Clínica Santa Ana (ficticia)").json()
    c.post(f"/api/pagos/{rep['reporte_id']}/aplicar")
    # «Paciente Uno» (ficticio) se pagó por otro lado (reporte de AP o a mano): el radicado sigue por cobrar
    d = c.get(f"/api/cirugias/{uno}").json()
    c.put(f"/api/cirugias/{uno}", json={**d, "estado": "PAGADA", "fecha_pago": "2026-02-10", "valor_honorarios": 500000})
    radicado = c.get("/api/inicio").json()["radicados"][0]
    assert (radicado["cirugias"], radicado["pagadas"], radicado["pendientes"]) == (2, 1, 2)   # dos y la línea sin cirugía
    lista = next(x for x in c.get("/api/pagos").json() if x["reporte_id"] == rep["reporte_id"])
    assert (lista["pagadas"], lista["pendientes"]) == (1, 2)

    # se paga la otra y se registra la tercera: ya no queda nada por cobrar en ese radicado
    d = c.get(f"/api/cirugias/{dos}").json()
    c.put(f"/api/cirugias/{dos}", json={**d, "estado": "PAGADA", "fecha_pago": "2026-02-11"})
    crear(c, "55555555", "2026-01-12", estado="PAGADA")
    assert c.get("/api/inicio").json()["radicados"] == []
    lista = next(x for x in c.get("/api/pagos").json() if x["reporte_id"] == rep["reporte_id"])
    assert (lista["estado"], lista["pagadas"], lista["pendientes"]) == ("APLICADO", 3, 0)
    # registrar después el pago de la clínica no cambia las fechas que ya tenían
    c.post(f"/api/pagos/{rep['reporte_id']}/pagar", json={"fecha_pago": "2026-03-20"})
    assert c.get(f"/api/cirugias/{uno}").json()["fecha_pago"] == "2026-02-10"
