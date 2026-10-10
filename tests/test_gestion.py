"""Detalle, corrección, archivo original, filtros, reportes y Excel."""
import io

from openpyxl import load_workbook

from tests.test_api import CONFIRMAR, cliente, subir  # noqa: F401  (fixture)
from tests.test_extraccion import foto


def orden(**cambios):
    return {**CONFIRMAR, **cambios}


def crear(cliente, **cambios):
    r = cliente.post("/api/cirugias", json=orden(**cambios))
    assert r.status_code == 201, r.text
    return r.json()["cirugia_id"]


def test_detalle_y_correccion(cliente):
    cid = crear(cliente, cirujano="Dr. Juan  Pérez", valor_honorarios="1.200.000")
    d = cliente.get(f"/api/cirugias/{cid}").json()
    assert d["cirujano"] == "JUAN PÉREZ" and d["valor_honorarios"] == 1200000
    assert d["paciente"]["numero_documento"] == "1020304050" and d["estado"] == "PROGRAMADA"

    corregida = orden(cirujano="JUAN PEREZ", estado="PAGADA", valor_honorarios=950000, fecha_pago="2026-11-30",
                      procedimientos=[{"codigo": "512101", "nombre": "Colecistectomía", "es_principal": True}],
                      aseguradora=None)
    assert cliente.put(f"/api/cirugias/{cid}", json=corregida).status_code == 200
    d = cliente.get(f"/api/cirugias/{cid}").json()
    assert d["estado"] == "PAGADA" and d["fecha_pago"] == "2026-11-30" and d["aseguradora"] is None
    assert [p["codigo"] for p in d["procedimientos"]] == ["512101"]


def test_no_se_ve_ni_corrige_cirugia_ajena(cliente):
    cid = crear(cliente)
    from app import auth
    cliente.cookies.set(auth.COOKIE, auth.crear_sesion(999))
    assert cliente.get(f"/api/cirugias/{cid}").status_code == 401  # usuario inexistente
    cliente.cookies.set(auth.COOKIE, auth.crear_sesion(1))
    assert cliente.get("/api/cirugias/99999").status_code == 404
    assert cliente.put("/api/cirugias/99999", json=orden()).status_code == 404


def test_archivo_original(cliente):
    datos = foto()
    doc = subir(cliente, datos).json()
    r = cliente.get(f"/api/documentos/{doc['documento_id']}/archivo")
    assert r.status_code == 200 and r.content == datos and r.headers["content-type"] == "image/jpeg"
    assert cliente.get("/api/documentos/99999/archivo").status_code == 404


def test_filtros(cliente):
    crear(cliente, cirujano="Ana Ruiz", fecha_procedimiento="2026-09-10")
    crear(cliente, cirujano="Luis Gómez", fecha_procedimiento="2026-10-05", aseguradora=None,
          paciente={"tipo_documento": "CC", "numero_documento": "555", "nombres": "Pedro", "apellidos": "Soto"})
    opciones = cliente.get("/api/opciones").json()
    ana = next(c["id"] for c in opciones["cirujanos"] if c["nombre"] == "ANA RUIZ")
    assert len(cliente.get(f"/api/cirugias?cirujano_id={ana}").json()) == 1
    assert len(cliente.get("/api/cirugias?desde=2026-10-01").json()) == 1
    assert cliente.get("/api/cirugias?aseguradora_id=PARTICULAR").json()[0]["paciente"] == "Pedro Soto"
    assert len(cliente.get("/api/cirugias?q=soto").json()) == 1
    assert cliente.get("/api/cirugias?aseguradora_id=x;drop").status_code == 422


def test_reporte_por_cirujano(cliente):
    crear(cliente, cirujano="Ana Ruiz", valor_honorarios=1000000, estado="PAGADA")
    crear(cliente, cirujano="Ana Ruiz", valor_honorarios=500000)
    crear(cliente, cirujano="Luis Gómez", valor_honorarios=300000)
    crear(cliente, cirujano="Luis Gómez", valor_honorarios=900000, estado="CANCELADA")
    crear(cliente, cirujano="Luis Gómez")  # sin valor
    r = cliente.get("/api/reportes").json()
    assert r["totales"]["cirugias"] == 4 and r["totales"]["canceladas"] == 1
    assert r["totales"]["total"] == 1800000 and r["totales"]["pagado"] == 1000000
    por = {x["nombre"]: x for x in r["por_cirujano"]}
    assert por["ANA RUIZ"]["total"] == 1500000 and por["ANA RUIZ"]["cirugias"] == 2
    assert por["LUIS GÓMEZ"]["cirugias"] == 2 and por["LUIS GÓMEZ"]["sin_valor"] == 1
    assert r["por_cirujano"][0]["nombre"] == "ANA RUIZ"  # el que más genera primero


def test_excel(cliente):
    crear(cliente, cirujano="Ana Ruiz", valor_honorarios=1000000)
    r = cliente.get("/api/cirugias.xlsx")
    assert r.status_code == 200 and "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Cirugías", "Por cirujano", "Por clínica", "Por aseguradora", "Por mes"]
    fila = [c.value for c in wb["Cirugías"][2]]
    assert "ANA RUIZ" in fila and 1000000 in fila
    assert wb["Por cirujano"]["A2"].value == "ANA RUIZ"
