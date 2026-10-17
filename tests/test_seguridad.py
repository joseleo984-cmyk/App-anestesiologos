"""Seguridad: cabeceras, CSRF, documentación cerrada, Excel sin fórmulas, tope de IA, auditoría y aislamiento
entre usuarios en todas las rutas con identificador (datos ficticios)."""
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import auth, config, exportar, extraccion, main, repositorio, seguridad, uso_ia
from app.modelos import ReportePagoExtraido
from tests.conftest import EMAIL, PASSWORD
from tests.test_api import CONFIRMAR
from tests.test_extraccion import foto
from tests.test_pagos import REPORTE
from tests.test_usuarios import app_multi, crear_cuenta, entrar, entrar_y_cambiar  # noqa: F401  (fixture)


def test_cabeceras_de_seguridad(app_multi):
    c = entrar(EMAIL, PASSWORD)
    for r in (c.get("/api/yo"), c.get("/login"), c.get("/salud")):
        h = r.headers
        assert h["x-content-type-options"] == "nosniff" and h["x-frame-options"] == "DENY"
        assert "frame-ancestors 'none'" in h["content-security-policy"] and "object-src 'none'" in h["content-security-policy"]
        assert "base-uri 'self'" in h["content-security-policy"] and h["referrer-policy"] == "same-origin"
    assert c.get("/api/yo").headers["cache-control"] == "no-store"      # datos de pacientes: sin cachés


def test_hsts_solo_con_https(app_multi, monkeypatch):
    c = TestClient(main.app)
    assert "strict-transport-security" not in c.get("/login").headers
    monkeypatch.setattr(config, "COOKIE_SEGURA", True)
    assert "max-age" in c.get("/login").headers["strict-transport-security"]


def test_la_documentacion_de_la_api_esta_cerrada(app_multi):
    c = TestClient(main.app)
    for ruta in ("/docs", "/redoc", "/openapi.json"):
        assert c.get(ruta).status_code == 404


def test_peticiones_de_otro_origen_se_rechazan(app_multi):
    c = entrar(EMAIL, PASSWORD)
    mala = {"Origin": "https://sitio-malo.example"}
    assert c.post("/api/logout", headers=mala).status_code == 403
    assert c.put("/api/cirugias/1", json=CONFIRMAR, headers={"Origin": "null"}).status_code == 403
    # el propio sitio y las peticiones sin Origin (apps, pruebas) pasan
    assert c.post("/api/logout", headers={"Origin": "http://testserver"}).status_code == 200
    # leer no cambia nada: una página ajena no puede leer la respuesta (CORS) ni lleva la cookie (SameSite)
    assert c.get("/api/yo", headers=mala).status_code in (200, 401)


def test_excel_no_ejecuta_formulas_de_los_datos(app_multi):
    assert seguridad.texto_seguro("=HYPERLINK(\"http://x\",\"a\")") == "'=HYPERLINK(\"http://x\",\"a\")"
    assert [seguridad.texto_seguro(v) for v in ("+1", "-2", "@x", "\tx", "Ana", None, 5)] == ["'+1", "'-2", "'@x", "'\tx", "Ana", None, 5]
    c = entrar(EMAIL, PASSWORD)
    datos = {**CONFIRMAR, "paciente": {**CONFIRMAR["paciente"], "nombres": "=1+1", "apellidos": "@SUM(A1)"},
             "cirujano": "-cmd|' /C calc'!A0"}
    assert c.post("/api/cirugias", json=datos).status_code == 201
    for ruta in ("/api/cirugias.xlsx", "/api/relacion.xlsx"):
        r = c.get(ruta)
        assert r.status_code == 200
        hoja = load_workbook(io.BytesIO(r.content)).active
        for fila in hoja.iter_rows():
            for celda in fila:
                if isinstance(celda.value, str) and celda.data_type == "f":
                    assert celda.value.startswith("=SUM(G"), f"fórmula inesperada en {ruta}: {celda.value}"   # solo el total propio


def test_tope_diario_de_lecturas_con_ia(app_multi, monkeypatch, conn):
    monkeypatch.setattr(config, "MAX_LECTURAS_IA_DIA", 2)
    c = entrar(EMAIL, PASSWORD)
    archivo = lambda n: {"archivo": ("orden.jpg", foto() + bytes([n]), "image/jpeg")}  # noqa: E731
    assert c.post("/api/documentos", files=archivo(1)).status_code == 201
    for _ in range(2):
        conn.execute("INSERT INTO uso_ia (anestesiologo_id, tipo, modelo) VALUES (1, 'orden', 'x')")
    conn.commit()
    r = c.post("/api/documentos", files=archivo(2))
    assert r.status_code == 429 and "límite" in r.json()["detail"]
    assert uso_ia.lecturas_ultimas_24h(conn, 1) == 2


def test_bitacora_de_auditoria(app_multi, conn):
    anonimo = TestClient(main.app)
    assert anonimo.post("/api/login", json={"email": "no-existe@clinica.co", "password": "x" * 12}).status_code == 401
    assert anonimo.post("/api/login", json={"email": EMAIL, "password": "equivocada-123"}).status_code == 401
    admin = entrar(EMAIL, PASSWORD)
    nueva = crear_cuenta(admin)
    admin.get("/api/cirugias.xlsx")
    nueva_clave = admin.post(f"/api/usuarios/{nueva['anestesiologo_id']}/restablecer").json()["clave_temporal"]
    admin.post("/api/copia-seguridad", json={"clave": "una-clave-larga-123"})
    eventos = [f["evento"] for f in admin.get("/api/auditoria").json()]
    for esperado in ("LOGIN_FALLO", "LOGIN_OK", "CUENTA_CREADA", "EXPORTAR_EXCEL", "CLAVE_RESTABLECIDA_ADMIN", "COPIA_SEGURIDAD"):
        assert esperado in eventos, esperado
    # la bitácora no guarda contraseñas ni datos de pacientes
    todo = str(conn.execute("SELECT * FROM auditoria").fetchall())
    assert "equivocada-123" not in todo and "clave_temporal" not in todo
    # solo el administrador la ve
    otro = entrar_y_cambiar("otro@clinica.co", nueva_clave)
    assert otro.get("/api/auditoria").status_code == 403
    assert anonimo.get("/api/auditoria").status_code == 401


def test_el_tiempo_de_login_no_delata_si_el_correo_existe(app_multi, monkeypatch):
    llamadas = []
    original = auth.verificar_password
    monkeypatch.setattr(auth, "verificar_password", lambda p, h: (llamadas.append(h), original(p, h))[1])
    c = TestClient(main.app)
    c.post("/api/login", json={"email": "nadie@clinica.co", "password": "x" * 12})
    assert llamadas == [auth.HASH_FALSO]       # con correo inexistente también se calcula un hash


def test_ninguna_ruta_con_identificador_deja_ver_ni_tocar_datos_de_otro(app_multi, monkeypatch):
    """Recorre las rutas de pagos, cirugías, documentos y finanzas con los identificadores de OTRO anestesiólogo."""
    monkeypatch.setattr(extraccion, "extraer_reporte_pago", lambda d, t: ReportePagoExtraido.model_validate(REPORTE))
    dueno = entrar(EMAIL, PASSWORD)
    otro = entrar_y_cambiar("otro@clinica.co", crear_cuenta(dueno)["clave_temporal"])

    doc = dueno.post("/api/documentos", files={"archivo": ("o.jpg", foto(), "image/jpeg")}).json()
    cid = dueno.post(f"/api/documentos/{doc['documento_id']}/confirmar", json={**CONFIRMAR, "numero_orden": "OR-1"}).json()["cirugia_id"]
    rep = dueno.post("/api/pagos", files={"archivo": ("r.pdf", b"%PDF-1.7 reporte", "application/pdf")}).json()
    rid, did = rep["reporte_id"], rep["filas"][0]["detalle_id"]
    meta = dueno.post("/api/finanzas/metas", json={"tipo": "VIAJE", "nombre": "Viaje", "valor_objetivo": 1000000}).json()
    partida = dueno.post("/api/finanzas/partidas", json={"clase": "GASTO", "categoria": "VIVIENDA", "nombre": "Arriendo", "valor_mensual": 500000})
    mov = dueno.post("/api/finanzas/movimientos", json={"clase": "APORTE", "meta_id": meta.get("meta_id"), "valor": 1000})

    def vacio(r):  # nada de lo del dueño en la respuesta
        return r.status_code in (403, 404, 409, 422) or r.json() in ([], {}, None)

    pruebas = [
        ("GET", f"/api/cirugias/{cid}", None), ("GET", f"/api/cirugias/{cid}/cobro", None), ("PUT", f"/api/cirugias/{cid}", CONFIRMAR),
        ("GET", f"/api/documentos/{doc['documento_id']}", None), ("GET", f"/api/documentos/{doc['documento_id']}/archivo", None),
        ("POST", f"/api/documentos/{doc['documento_id']}/confirmar", CONFIRMAR),
        ("GET", f"/api/pagos/{rid}", None), ("GET", f"/api/pagos/{rid}/archivo", None), ("GET", f"/api/pagos/{rid}/liquidacion", None),
        ("POST", f"/api/pagos/{rid}/aplicar", None), ("POST", f"/api/pagos/{rid}/cruzar", None),
        ("POST", f"/api/pagos/{rid}/crear-faltantes", None), ("POST", f"/api/pagos/{rid}/pagar", {"fecha_pago": "2026-03-20"}),
        ("POST", f"/api/pagos/{rid}/pagar-pacientes", {"fecha_pago": "2026-03-20", "lineas": [{"detalle_id": did, "valor": 1000}]}),
        ("DELETE", f"/api/pagos/{rid}", None),
        ("GET", f"/api/pagos/lineas/{did}/candidatas", None),
        ("POST", f"/api/pagos/lineas/{did}/vincular", {"cirugia_id": cid, "cedula_correcta": "app"}),
        ("POST", f"/api/pagos/lineas/{did}/radicado", {"radicado_id": rid}),
    ]
    if meta.get("meta_id"):
        pruebas += [("PUT", f"/api/finanzas/metas/{meta['meta_id']}", {"tipo": "VIAJE", "nombre": "Mío", "valor_objetivo": 5}),
                    ("DELETE", f"/api/finanzas/metas/{meta['meta_id']}", None)]
    if partida.status_code == 201:
        pruebas += [("DELETE", f"/api/finanzas/partidas/{partida.json()['partida_id']}", None)]
    if mov.status_code == 201:
        pruebas += [("DELETE", f"/api/finanzas/movimientos/{mov.json()['movimiento_id']}", None)]
    for metodo, ruta, cuerpo in pruebas:
        r = otro.request(metodo, ruta, json=cuerpo) if cuerpo is not None else otro.request(metodo, ruta)
        assert r.status_code < 500, (metodo, ruta, r.status_code)
        assert r.status_code in (403, 404, 409, 422), f"{metodo} {ruta} respondió {r.status_code}: otro usuario no debería poder"
    # lo del dueño sigue intacto
    assert dueno.get(f"/api/cirugias/{cid}").status_code == 200 and dueno.get(f"/api/pagos/{rid}").status_code == 200
    assert otro.get("/api/pagos").json() == [] and otro.get("/api/cirugias").json() == []
