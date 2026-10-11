"""Relación de órdenes: filtros de selección múltiple, nº de orden y Excel."""
import io

from openpyxl import load_workbook

from tests.test_api import CONFIRMAR, cliente  # noqa: F401  (fixture)

SANTA_ANA = {"nombre": "Clínica Santa Ana de Dios S.A.S.", "ciudad": "Barranquilla"}
CARIBE = {"nombre": "Clínica del Caribe", "ciudad": "Barranquilla"}


def cirugia(cliente, doc, aseguradora, clinica, valor=None, estado="REALIZADA", orden=None, fecha="2026-10-15"):
    datos = {**CONFIRMAR, "paciente": {**CONFIRMAR["paciente"], "numero_documento": doc},
             "aseguradora": aseguradora, "clinica": clinica, "estado": estado, "valor_honorarios": valor,
             "numero_orden": orden, "fecha_procedimiento": fecha}
    r = cliente.post("/api/cirugias", json=datos)
    assert r.status_code == 201, r.text
    return r.json()["cirugia_id"]


def docs(filas):
    return sorted(f["cedula"] for f in filas)


def preparar(cliente):
    cirugia(cliente, "1001", {"nombre": "Colsanitas", "tipo": "PREPAGADA"}, CARIBE, "649.405", orden="2737/8784")
    cirugia(cliente, "1002", {"nombre": "Allianz", "tipo": "PREPAGADA"}, SANTA_ANA, 893269)
    cirugia(cliente, "1003", {"nombre": "Sura", "tipo": "EPS"}, SANTA_ANA, 100000)
    cirugia(cliente, "1004", {"nombre": "Sura", "tipo": "EPS"}, CARIBE, 50000)
    cirugia(cliente, "1005", None, CARIBE)                                         # particular
    cirugia(cliente, "1006", {"nombre": "Allianz", "tipo": "PREPAGADA"}, SANTA_ANA, 1, estado="CANCELADA")
    cirugia(cliente, "1007", {"nombre": "Allianz", "tipo": "PREPAGADA"}, CARIBE, 2, estado="PAGADA")


def test_filtros_de_seleccion_multiple(cliente):
    preparar(cliente)
    rel = lambda **p: cliente.get("/api/relacion", params=p).json()  # noqa: E731
    assert docs(rel()) == ["1001", "1002", "1003", "1004", "1005", "1007"]           # canceladas nunca
    assert docs(rel(tipo="PREPAGADA")) == ["1001", "1002", "1007"]
    assert docs(rel(tipo=["EPS", "PARTICULAR"])) == ["1003", "1004", "1005"]
    opciones = cliente.get("/api/opciones").json()
    ids = {c["nombre"]: c["id"] for c in opciones["clinicas"]}
    santa_ana = ids["CLÍNICA SANTA ANA DE DIOS S.A.S."]
    assert docs(rel(clinica_id=santa_ana)) == ["1002", "1003"]
    assert docs(rel(clinica_id=list(ids.values()))) == ["1001", "1002", "1003", "1004", "1005", "1007"]
    aseg = {a["nombre"]: a["id"] for a in opciones["aseguradoras"]}
    assert docs(rel(aseguradora_id=[aseg["COLSANITAS"], aseg["SURA"]])) == ["1001", "1003", "1004"]
    assert docs(rel(tipo="PREPAGADA", clinica_id=santa_ana)) == ["1002"]
    assert docs(rel(tipo="PREPAGADA", estado="PENDIENTE")) == ["1001", "1002"]      # sin pagar
    assert cliente.get("/api/relacion", params={"tipo": "OTRO"}).status_code == 422


def test_fila_con_la_estructura_del_formato(cliente):
    preparar(cliente)
    fila = next(f for f in cliente.get("/api/relacion").json() if f["cedula"] == "1001")
    assert fila["entidad"] == "COLSANITAS" and fila["paciente"] == "María Gómez"
    assert fila["procedimiento"] == "Apendicectomía" and fila["numero_orden"] == "2737/8784"
    assert float(fila["valor"]) == 649405


def test_numero_de_orden_se_guarda_y_corrige(cliente):
    cid = cirugia(cliente, "2001", {"nombre": "Allianz", "tipo": "PREPAGADA"}, SANTA_ANA, orden="A-1")
    assert cliente.get(f"/api/cirugias/{cid}").json()["numero_orden"] == "A-1"
    datos = {**CONFIRMAR, "numero_orden": "B-2"}
    assert cliente.put(f"/api/cirugias/{cid}", json=datos).status_code == 200
    assert cliente.get(f"/api/cirugias/{cid}").json()["numero_orden"] == "B-2"


def test_excel_solo_filas_elegidas(cliente):
    preparar(cliente)
    filas = cliente.get("/api/relacion", params={"tipo": "PREPAGADA", "estado": "PENDIENTE"}).json()
    r = cliente.get("/api/relacion.xlsx", params={"tipo": "PREPAGADA", "estado": "PENDIENTE",
                                                  "ids": [f["cirugia_id"] for f in filas], "mes": "Octubre 2026"})
    assert r.status_code == 200
    ws = load_workbook(io.BytesIO(r.content)).active
    valores = [[c.value for c in fila] for fila in ws.iter_rows()]
    assert "MES A FACTURAR: Octubre 2026" in valores[1]
    assert valores[4] == ["#", "ENTIDAD", "PACIENTE", "CÉDULA", "PROCEDIMIENTO", "Nº ORDEN", "VALOR"]
    assert [v[3] for v in valores[5:7]] == ["1001", "1002"]
    assert valores[7][5:] == ["TOTAL", "=SUM(G6:G7)"]
    # Solo una fila elegida
    r = cliente.get("/api/relacion.xlsx", params={"tipo": "PREPAGADA", "ids": [filas[1]["cirugia_id"]]})
    ws = load_workbook(io.BytesIO(r.content)).active
    assert [c.value for c in ws[6]][3] == "1002" and ws.cell(7, 6).value == "TOTAL"
